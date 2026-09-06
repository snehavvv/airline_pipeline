"""
ASG Airlines — Documentation Word (.docx) Generator
Generates publication-quality .docx document deliverable for ASG Airlines.
"""

from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build_docx_documentation(output_file: Path, img_dir: Path):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    def set_cell_background(cell, fill_hex):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    # Title Banner
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("✈️ ASG AIRLINES — DATA ENGINEERING PIPELINE\n& POWER BI REPORTING SUITE")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Full Technical Documentation, End-to-End Architecture Walkthrough & Data Governance Report")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph()

    # Section 1: Executive Overview
    h1 = doc.add_heading("1. Executive Overview & Deliverables Inventory", level=1)
    h1.runs[0].font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph("This enterprise documentation outlines the end-to-end local data engineering pipeline and Power BI reporting solution developed for ASG Airlines. The solution ingests, cleanses, standardizes, transforms, governs, and models data across all 4 core operational tables: flights, bookings, passengers, and payments.")

    # Table of Deliverables
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ["Deliverable Item", "Repository Location", "Format", "Status"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "1E3A8A")

    deliverables = [
        ("Power BI Report File", "powerbi/ASG_Airlines_Dashboard.pbix", ".pbix (373 KB)", "✅ Built & Included"),
        ("Analytical Report Previews", "powerbi/previews/", ".png (4 page renders)", "✅ Built & Included"),
        ("Architecture & Data Flow Diagram", "docs/architecture_data_flow_diagram.png", ".png Visual Diagram", "✅ Built & Included"),
        ("Relational ERD Data Model", "docs/relational_data_model_diagram.png", ".png Visual ERD", "✅ Built & Included"),
        ("Solution Architecture Walkthrough", "Solution_Walkthrough.md", "Narrative Walkthrough", "✅ Built & Included"),
        ("Enterprise ETL Pipeline Script", "run_pipeline.py", "Python 3.10 Script", "✅ Fail-safe Try/Except"),
        ("Sample Raw Dataset", "data/raw/sample_usecase_airlines.xlsx", "Excel Workbook", "✅ Out-of-the-box Ready"),
        ("Cleaned & Masked Datasets", "data/cleaned/*.csv", "5 CSV Files (Unified Master)", "✅ 100% Table Joined"),
        ("KPI Aggregations Engine", "data/aggregated/*.csv", "12 CSV Tables (7 Bonus KPIs)", "✅ 12 Aggregations")
    ]

    for item, loc, fmt, stat in deliverables:
        row_cells = table.add_row().cells
        row_cells[0].text = item
        row_cells[1].text = loc
        row_cells[2].text = fmt
        row_cells[3].text = stat
        for cell in row_cells:
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph()

    # Section 2: Architecture Diagram
    h2 = doc.add_heading("2. Pipeline Architecture & Data Flow Diagram", level=1)
    h2.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph("The pipeline operates across 6 modular stages: Raw Data Layer -> Ingestion & Validation -> Cleaning & Transformation -> PII Governance -> KPI Aggregation -> Power BI & Export Layer.")

    arch_img = img_dir / "architecture_data_flow_diagram.png"
    if arch_img.exists():
        doc.add_picture(str(arch_img), width=Inches(6.5))
        p_cap = doc.add_paragraph("Figure 1: End-to-End Enterprise Data Pipeline Architecture & Data Flow Diagram")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(9)
        p_cap.runs[0].font.italic = True

    doc.add_paragraph()

    # Section 3: Data Model ERD Diagram
    h3 = doc.add_heading("3. Relational Data Model & ERD Schema", level=1)
    h3.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph("The relational star schema unifies all 4 core entities (passengers, bookings, flights, payments) into a master analytical fact table (master_dataset.csv).")

    erd_img = img_dir / "relational_data_model_diagram.png"
    if erd_img.exists():
        doc.add_picture(str(erd_img), width=Inches(6.5))
        p_cap2 = doc.add_paragraph("Figure 2: ASG Airlines Relational Data Model & Entity-Relationship Schema")
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.runs[0].font.size = Pt(9)
        p_cap2.runs[0].font.italic = True

    doc.add_paragraph()

    # Section 4: Power BI Report
    h4 = doc.add_heading("4. Power BI Interactive Report & Analytical Previews", level=1)
    h4.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph("The deliverables include ASG_Airlines_Dashboard.pbix containing 4 visual dashboard pages, data models, and DAX measures. The preview below highlights the Operations Overview analytical page.")

    pbi_img = img_dir.parent / "powerbi" / "previews" / "asg_airlines_dashboard_preview.png"
    if pbi_img.exists():
        doc.add_picture(str(pbi_img.resolve()), width=Inches(6.5))
        p_cap3 = doc.add_paragraph("Figure 3: Operations Overview Analytical Report Preview")
        p_cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap3.runs[0].font.size = Pt(9)
        p_cap3.runs[0].font.italic = True

    output_file.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_file)
    print(f"Created Word documentation (.docx): {output_file}")

if __name__ == "__main__":
    repo_root = Path(__file__).resolve().parent.parent
    docs_dir = repo_root / "docs"
    build_docx_documentation(repo_root / "ASG_Airlines_Documentation.docx", docs_dir)
    
    asg_root = repo_root.parent / "ASG-Airlines-Pipeline"
    if asg_root.exists():
        build_docx_documentation(asg_root / "ASG_Airlines_Documentation.docx", asg_root / "docs")
